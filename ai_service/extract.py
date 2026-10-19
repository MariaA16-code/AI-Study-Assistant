"""Read PDF / DOCX / TXT into pages with page numbers and chapter hints."""
import hashlib
import io
import logging
import re
import threading
from dataclasses import dataclass, field
from pathlib import Path

from . import settings
from .vision import describe_many, figure_lines, vision_available

logger = logging.getLogger(__name__)

MIN_TEXT_CHARS = 20  # a PDF page with less text than this is treated as a scanned image
PAGE_CHARS = 3000  # DOCX/TXT have no real pages, so we make ~3000-char sections
CHAPTER_RE = re.compile(r"^(chapter|unit|lesson)\s+(\d+|[ivxlc]+)\b", re.IGNORECASE)
# Table-of-contents entries: dot leaders, or a chapter/section entry ending in a page number
TOC_RE = re.compile(
    r"([.…·]\s?){4,}"  # leaders like ...... or ………
    r"|^\s*((chapter|unit|lesson)\s+\w+|\d+(\.\d+)+)\b.*(\t| {2,})\s*(\d+|[ivxlc]+)\s*$"
    r"|^[^\t]{2,80}\t\s*(\d{1,3}|[ivxlc]{1,6})\s*$",  # short line, tab, page number
    re.IGNORECASE,
)
_BARE_HEADING_RE = re.compile(r"(chapter|unit|lesson)\s+(\d+|[ivxlc]+)\.?", re.IGNORECASE)


def _is_chapter_line(line: str) -> bool:
    line = line.strip()
    if not CHAPTER_RE.match(line) or TOC_RE.search(line):
        return False
    if _BARE_HEADING_RE.fullmatch(line):  # just "Chapter 3" or "Chapter 3."
        return True
    # a heading is short and does not read like a sentence
    return len(line) <= 80 and not line.endswith((".", "!", "?", ":"))


@dataclass
class Page:
    number: int  # 1-based page (or section) number
    text: str
    chapter: str | None = None


@dataclass
class ExtractedDocument:
    pages: list[Page]
    warnings: list[str] = field(default_factory=list)


def _detect_chapter(text: str) -> str | None:
    for line in text.splitlines()[:8]:  # headings sit near the top of a page
        line = line.strip()
        if _is_chapter_line(line):
            return line[:80]
    return None


_ocr_lock = threading.Lock()
_ocr_engine = None


def _get_ocr_engine():
    """Load the OCR engine once. Raises ImportError when rapidocr is not installed."""
    global _ocr_engine
    with _ocr_lock:
        if _ocr_engine is None:
            from rapidocr import RapidOCR

            _ocr_engine = RapidOCR(params={"Global.log_level": "error"})
        return _ocr_engine


def _boxes_to_text(items: list[tuple[float, float, float, float, str]]) -> str:
    """Turn OCR boxes (x0, y0, x1, y1, text) into lines of text.

    Boxes on the same line are joined in reading order. A wide gap between two boxes
    (a table column) is written as ' | ', so table rows keep their cells apart.
    """
    if not items:
        return ""
    heights = sorted(y1 - y0 for _, y0, _, y1, _ in items)
    median = heights[len(heights) // 2] or 1.0

    def centre(item):
        return (item[1] + item[3]) / 2

    rows: list[list] = []
    current: list = []
    for item in sorted(items, key=centre):
        if current and abs(centre(item) - sum(centre(c) for c in current) / len(current)) > 0.6 * median:
            rows.append(current)
            current = []
        current.append(item)
    if current:
        rows.append(current)

    lines = []
    for row in rows:
        row.sort(key=lambda it: it[0])
        parts = [row[0][4]]
        for prev, nxt in zip(row, row[1:]):
            parts.append(" | " if nxt[0] - prev[2] > 2 * median else " ")
            parts.append(nxt[4])
        line = "".join(parts).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def _ocr_image(image) -> str:
    """Read the text in a PIL image with OCR. Raises ImportError when rapidocr is not installed."""
    import numpy as np

    engine = _get_ocr_engine()
    result = engine(np.array(image.convert("RGB")))
    if result is None or result.boxes is None or result.txts is None:
        return ""
    items = []
    for box, text, score in zip(result.boxes, result.txts, result.scores):
        if float(score) < 0.5 or not str(text).strip():  # skip uncertain noise
            continue
        xs = [pt[0] for pt in box]
        ys = [pt[1] for pt in box]
        items.append((min(xs), min(ys), max(xs), max(ys), str(text).strip()))
    return _boxes_to_text(items)


def _ocr_page_text(page) -> str:
    """Read a scanned PDF page with OCR."""
    return _ocr_image(page.to_image(resolution=settings.OCR_DPI).original)


FIG_MIN_SIDE = 100  # points; an embedded picture smaller than this is an icon or bullet
FIG_MIN_AREA = 0.04  # share of the page an embedded picture must cover
FIG_VECTOR_MIN = 4  # boxes + curves on a page that make it a drawn diagram or chart
FIG_REPEAT_PAGES = 3  # a picture on this many pages (and over 30% of them) is a logo or background


def _picture_ids(page) -> set:
    """Ids of the large embedded pictures on a page (small ones are icons)."""
    area = float(page.width * page.height) or 1.0
    found = set()
    for img in page.images:
        w = abs(float(img["x1"]) - float(img["x0"]))
        h = abs(float(img["bottom"]) - float(img["top"]))
        if w < FIG_MIN_SIDE or h < FIG_MIN_SIDE or (w * h) / area < FIG_MIN_AREA:
            continue
        stream = img.get("stream")
        ident = getattr(stream, "objid", None)
        found.add(ident if ident is not None else (round(w), round(h), img.get("name")))
    return found


def _has_drawn_figure(page) -> bool:
    """Boxes and curves drawn on the page: a flowchart, UML box, bar chart or similar."""
    area = float(page.width * page.height) or 1.0
    shapes = len(page.curves)
    for rect in page.rects:
        w, h = float(rect["width"]), float(rect["height"])
        if w < 8 or h < 8 or (w * h) / area > 0.9:  # table rules, thin lines, page background
            continue
        shapes += 1
    return shapes >= FIG_VECTOR_MIN


def _read_pdf(path: Path) -> tuple[list[Page], list[str]]:
    import pdfplumber

    pages: list[Page] = []
    empty = ocr_used = ocr_skipped = 0
    ocr_missing = False
    fig_used = fig_failed = fig_skipped = 0
    with pdfplumber.open(str(path)) as pdf:
        picture_ids: list[set] = []
        drawn: list[bool] = []
        for i, p in enumerate(pdf.pages, start=1):
            text = (p.extract_text() or "").strip()
            scanned = len(text) < MIN_TEXT_CHARS
            if scanned and settings.OCR_ENABLED and not ocr_missing:
                if ocr_used >= settings.OCR_MAX_PAGES:
                    ocr_skipped += 1
                else:
                    try:
                        ocr_text = _ocr_page_text(p)
                        if len(ocr_text) > len(text):
                            text = ocr_text
                            ocr_used += 1
                    except ImportError:
                        ocr_missing = True
                    except Exception:
                        logger.exception("OCR failed on page %s of %s", i, path.name)
            if not text:
                empty += 1
            pages.append(Page(i, text))
            # pages without a text layer are read by OCR only; vision looks at the others
            if scanned or not vision_available():
                picture_ids.append(set())
                drawn.append(False)
            else:
                try:
                    picture_ids.append(_picture_ids(p))
                    drawn.append(_has_drawn_figure(p))
                except Exception:
                    logger.warning("Could not look for figures on page %s", i, exc_info=True)
                    picture_ids.append(set())
                    drawn.append(False)

        if vision_available():
            counts: dict = {}
            for ids in picture_ids:
                for ident in ids:
                    counts[ident] = counts.get(ident, 0) + 1
            repeated = {
                ident
                for ident, n in counts.items()
                if n >= FIG_REPEAT_PAGES and n > 0.3 * len(pages)
            }
            flagged = [
                idx for idx in range(len(pages))
                if drawn[idx] or (picture_ids[idx] - repeated)
            ]
            fig_skipped = max(0, len(flagged) - settings.VISION_MAX_PAGES)
            flagged = flagged[: settings.VISION_MAX_PAGES]
            images = []
            for idx in flagged:
                try:
                    images.append(pdf.pages[idx].to_image(resolution=settings.VISION_DPI).original)
                except Exception:
                    logger.warning("Could not render page %s for vision", idx + 1, exc_info=True)
                    images.append(None)
            todo = [(idx, img) for idx, img in zip(flagged, images) if img is not None]
            fig_failed += len(flagged) - len(todo)
            results = describe_many([img for _, img in todo], page=True)
            for (idx, _), description in zip(todo, results):
                if description is None:
                    fig_failed += 1
                elif description:
                    lines = figure_lines(description)
                    pages[idx].text = (pages[idx].text + "\n" + "\n".join(lines)).strip()
                    fig_used += 1

    current = None
    for page in pages:
        found = _detect_chapter(page.text)
        if found:
            current = found
        page.chapter = current

    warnings = []
    if ocr_used:
        warnings.append(
            f"{ocr_used} page(s) had no text layer and were read with OCR. "
            "OCR can make small mistakes, especially with tables and handwriting."
        )
    if ocr_missing:
        warnings.append("Some pages are scanned images, but OCR is not installed (pip install rapidocr).")
    if ocr_skipped:
        warnings.append(
            f"OCR was limited to {settings.OCR_MAX_PAGES} pages, so {ocr_skipped} scanned page(s) were skipped."
        )
    if fig_used:
        warnings.append(
            f"{fig_used} page(s) with diagrams, charts or pictures were described with AI vision. "
            "Descriptions can contain mistakes, especially arrow directions."
        )
    if fig_failed:
        warnings.append(f"AI vision could not describe {fig_failed} page(s) with figures.")
    if fig_skipped:
        warnings.append(
            f"AI vision was limited to {settings.VISION_MAX_PAGES} pages, "
            f"so {fig_skipped} page(s) with figures were not described."
        )
    if pages and empty / len(pages) > 0.5:
        warnings.append(
            "Most pages have no extractable text. This looks like a scanned PDF, "
            "and OCR could not read it."
        )
    return pages, warnings


def _paragraphs_to_pages(items: list[tuple[str, bool]]) -> list[Page]:
    """items: (text, is_heading). Groups paragraphs into ~PAGE_CHARS sections."""
    pages: list[Page] = []
    buf: list[str] = []
    size = 0
    current = None
    page_chapter = None

    def flush():
        nonlocal buf, size
        if buf:
            pages.append(Page(len(pages) + 1, "\n".join(buf), page_chapter))
        buf, size = [], 0

    for text, is_heading in items:
        if (is_heading and not TOC_RE.search(text)) or _is_chapter_line(text):
            flush()
            current = text[:80]
        if not buf:
            page_chapter = current
        buf.append(text)
        size += len(text)
        if size >= PAGE_CHARS:
            flush()
    flush()
    return pages


def _docx_blocks(document):
    """Paragraphs and tables in the order they appear in the document."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


class _DocxImages:
    """Reads the text inside pictures of a DOCX with OCR. Each distinct picture is read once."""

    MIN_SIDE = 60  # smaller pictures are icons or bullets, not worth reading

    def __init__(self, document):
        self.document = document
        self.cache: dict[str, list[str]] = {}
        self.attempts = 0
        self.with_text = 0
        self.skipped = 0
        self.missing = False
        self.vision_used = 0
        self.vision_failed = 0
        self.vision_skipped = 0

    def _open(self, blob: bytes):
        """The picture as a PIL image, or None when it is too small or cannot be opened."""
        from PIL import Image

        try:
            image = Image.open(io.BytesIO(blob))
            image.load()
        except Exception:  # unsupported picture formats such as EMF or WMF
            return None
        return None if min(image.size) < self.MIN_SIDE else image

    def prefetch(self) -> None:
        """Describe all pictures with vision first (a few at a time). Failures fall back to OCR."""
        if not vision_available():
            return
        keys, images = [], []
        seen = set()
        for rid in self.document.element.body.xpath(".//a:blip/@r:embed"):
            part = self.document.part.related_parts.get(rid)
            if part is None:
                continue
            key = hashlib.md5(part.blob).hexdigest()
            if key in seen:
                continue
            seen.add(key)
            image = self._open(part.blob)
            if image is None:
                continue
            if len(images) >= settings.VISION_MAX_IMAGES:
                self.vision_skipped += 1
                continue
            keys.append(key)
            images.append(image)
        for key, description in zip(keys, describe_many(images)):
            if description is None:
                self.vision_failed += 1  # not cached, so lines() reads it with OCR
                continue
            lines = figure_lines(description)
            self.cache[key] = lines
            if lines:
                self.vision_used += 1

    def lines(self, paragraph) -> list[str]:
        """Text lines found in the pictures of one paragraph (empty when there are none)."""
        found = []
        for rid in paragraph._p.xpath(".//a:blip/@r:embed"):
            part = self.document.part.related_parts.get(rid)
            if part is None:
                continue
            key = hashlib.md5(part.blob).hexdigest()
            if key not in self.cache:
                if not settings.OCR_ENABLED or self.missing:
                    continue  # not described by vision and OCR is unavailable
                self.cache[key] = self._read(part.blob)
            found.extend(self.cache[key])
        return found

    def _read(self, blob: bytes) -> list[str]:
        from PIL import Image

        try:
            image = Image.open(io.BytesIO(blob))
            if min(image.size) < self.MIN_SIDE:
                return []
            if self.attempts >= settings.OCR_MAX_IMAGES:
                self.skipped += 1
                return []
            self.attempts += 1
            text = _ocr_image(image)
        except ImportError:
            self.missing = True
            return []
        except Exception:  # unsupported picture formats such as EMF or WMF
            logger.warning("Could not read a picture inside the DOCX", exc_info=True)
            return []
        lines = []
        for line in text.splitlines():
            if CHAPTER_RE.match(line):  # a line inside a picture must never start a new chapter
                line = "- " + line
            lines.append(line)
        if lines:
            self.with_text += 1
        return lines


def _table_rows(table, images: "_DocxImages | None" = None) -> list[str]:
    """One text line per table row, cells joined with ' | ' (merged cells are counted once).

    Tables nested inside a cell, and text inside pictures in a cell, are read right after the row.
    """
    rows = []
    for row in table.rows:
        cells, seen, nested, extra = [], set(), [], []
        for cell in row.cells:
            if id(cell._tc) in seen:  # a merged cell shows up once per column it spans
                continue
            seen.add(id(cell._tc))
            text = " ".join(cell.text.split())
            if text:
                cells.append(text)
            nested.extend(cell.tables)
            if images is not None:
                for paragraph in cell.paragraphs:
                    extra.extend(images.lines(paragraph))
        line = " | ".join(cells)
        if line:
            if CHAPTER_RE.match(line):  # a table row must never start a new chapter
                line = "- " + line
            rows.append(line)
        rows.extend(extra)
        for inner in nested:
            rows.extend(_table_rows(inner, images))
    return rows


def _read_docx(path: Path) -> tuple[list[Page], list[str]]:
    import docx
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    images = _DocxImages(document)
    images.prefetch()
    items = []
    for block in _docx_blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                style = block.style.name if block.style is not None else ""
                items.append((text, style.startswith("Heading 1")))
            items.extend((line, False) for line in images.lines(block))
        else:
            items.extend((row, False) for row in _table_rows(block, images))

    warnings = ["DOCX has no fixed pages; page numbers are approximate sections."]
    if images.vision_used:
        warnings.append(
            f"{images.vision_used} picture(s) in the DOCX were described with AI vision. "
            "Descriptions can contain mistakes, especially arrow directions."
        )
    if images.vision_failed:
        warnings.append(
            f"AI vision could not describe {images.vision_failed} picture(s); "
            "they were read with OCR when possible."
        )
    if images.vision_skipped:
        warnings.append(
            f"AI vision was limited to {settings.VISION_MAX_IMAGES} pictures, "
            f"so {images.vision_skipped} picture(s) were read with OCR only."
        )
    if images.with_text:
        warnings.append(
            f"{images.with_text} picture(s) in the DOCX were read with OCR. "
            "OCR can make small mistakes."
        )
    if images.missing:
        warnings.append(
            "The DOCX has pictures, but OCR is not installed, so text inside pictures "
            "is not read (pip install rapidocr)."
        )
    if images.skipped:
        warnings.append(
            f"OCR was limited to {settings.OCR_MAX_IMAGES} pictures, "
            f"so {images.skipped} picture(s) were skipped."
        )
    return _paragraphs_to_pages(items), warnings


def _read_txt(path: Path) -> tuple[list[Page], list[str]]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    # one item per line, so headings are found even without blank lines around them
    items = [(line.strip(), False) for line in raw.splitlines() if line.strip()]
    warnings = ["TXT has no fixed pages; page numbers are approximate sections."]
    return _paragraphs_to_pages(items), warnings


def extract_document(file_path: str) -> ExtractedDocument:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = path.suffix.lower()
    if ext == ".pdf":
        pages, warnings = _read_pdf(path)
    elif ext == ".docx":
        pages, warnings = _read_docx(path)
    elif ext == ".txt":
        pages, warnings = _read_txt(path)
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT.")

    if not any(p.text.strip() for p in pages):
        warnings.append("No readable text found in this file.")
    return ExtractedDocument(pages, warnings)