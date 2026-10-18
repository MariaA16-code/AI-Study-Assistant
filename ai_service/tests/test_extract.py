import docx
import pytest

from ai_service.extract import extract_document


def test_txt_pages_and_chapters(tmp_path):
    f = tmp_path / "book.txt"
    f.write_text(
        "Chapter 1 Cells\n\nCells are the unit of life.\n\n"
        "Chapter 2 Energy\n\nPlants make food.",
        encoding="utf-8",
    )
    doc = extract_document(str(f))
    assert doc.pages[0].chapter.startswith("Chapter 1")
    assert any(p.chapter and p.chapter.startswith("Chapter 2") for p in doc.pages)
    assert "Cells are the unit of life." in doc.pages[0].text


def test_docx_heading_becomes_chapter(tmp_path):
    d = docx.Document()
    d.add_heading("Photosynthesis", level=1)
    d.add_paragraph("Plants make food using light.")
    path = tmp_path / "notes.docx"
    d.save(str(path))
    doc = extract_document(str(path))
    assert doc.pages[0].chapter == "Photosynthesis"


def test_unsupported_type(tmp_path):
    f = tmp_path / "a.xyz"
    f.write_text("x")
    with pytest.raises(ValueError):
        extract_document(str(f))


def test_missing_file():
    with pytest.raises(FileNotFoundError):
        extract_document("nope.pdf")

def test_toc_lines_are_not_chapters(tmp_path):
    f = tmp_path / "toc.txt"
    f.write_text(
        "Chapter 1 Intro ........ 5\n\nChapter 1 Intro\n\nReal text here.",
        encoding="utf-8",
    )
    doc = extract_document(str(f))
    assert all("...." not in (p.chapter or "") for p in doc.pages)
    assert doc.pages[-1].chapter == "Chapter 1 Intro"    

def test_docx_toc_entry_with_tab_is_not_a_chapter(tmp_path):
    d = docx.Document()
    d.add_paragraph("Chapter 1 Intro\t5")
    d.add_heading("Chapter 1 Intro", level=1)
    d.add_paragraph("Real content about the introduction of the project goes here.")
    path = tmp_path / "toc.docx"
    d.save(str(path))
    doc = extract_document(str(path))
    assert doc.pages[-1].chapter == "Chapter 1 Intro"
    assert all("\t" not in (p.chapter or "") for p in doc.pages)
def test_txt_heading_without_blank_lines(tmp_path):
    f = tmp_path / "thesis.txt"
    f.write_text(
        "Chapter 1. Introduction\n1.1 Overview\nShieldNet is a platform.\n"
        "Chapter 5 describes the system testing process.\n"
        "Chapter 2. Literature Review\n2.1 Related work\nOther systems exist.",
        encoding="utf-8",
    )
    doc = extract_document(str(f))
    chapters = [p.chapter for p in doc.pages if p.chapter]
    assert chapters == ["Chapter 1. Introduction", "Chapter 2. Literature Review"]

def test_txt_toc_with_spaces_is_not_a_chapter(tmp_path):
    f = tmp_path / "toc.txt"
    f.write_text(
        "Table of contents\n"
        "Chapter 1. Introduction          1\n"
        "Chapter 2. Methods          14\n\n"
        "Chapter 1. Introduction\nReal text.\n",
        encoding="utf-8",
    )
    doc = extract_document(str(f))
    assert [p.chapter for p in doc.pages] == [None, "Chapter 1. Introduction"]  

def test_docx_table_text_is_read_in_order(tmp_path):
       d = docx.Document()
       d.add_paragraph("Chapter 1. Testing", style="Heading 1")
       d.add_paragraph("Before the table.")
       table = d.add_table(rows=2, cols=3)
       table.cell(0, 0).text = "Test case"
       table.cell(0, 1).text = "Expected"
       table.cell(0, 2).text = "Result"
       table.cell(1, 0).text = "Login with a wrong password"
       table.cell(1, 1).text = "Error message"
       table.cell(1, 2).text = "Pass"
       d.add_paragraph("After the table.")
       path = tmp_path / "table.docx"
       d.save(path)

       doc = extract_document(str(path))
       text = "\n".join(p.text for p in doc.pages)

       assert "Login with a wrong password | Error message | Pass" in text
       assert text.index("Before the table.") < text.index("Test case") < text.index("After the table.")
       assert not any("Tables in the DOCX" in w for w in doc.warnings)


def test_docx_merged_table_cell_is_read_once(tmp_path):
       d = docx.Document()
       table = d.add_table(rows=1, cols=2)
       merged = table.cell(0, 0).merge(table.cell(0, 1))
       merged.text = "Merged heading"
       path = tmp_path / "merged.docx"
       d.save(path)

       text = "\n".join(p.text for p in extract_document(str(path)).pages)
       assert text.count("Merged heading") == 1


def test_docx_table_row_never_starts_a_chapter(tmp_path):
       d = docx.Document()
       d.add_paragraph("Chapter 1. Intro", style="Heading 1")
       table = d.add_table(rows=1, cols=2)
       table.cell(0, 0).text = "Chapter 9"
       table.cell(0, 1).text = "see appendix"
       path = tmp_path / "row.docx"
       d.save(path)

       doc = extract_document(str(path))
       assert {p.chapter for p in doc.pages} == {"Chapter 1. Intro"}   

 # ---------- scanned PDFs (OCR) and nested tables ----------

from PIL import Image, ImageDraw, ImageFont

from ai_service import extract as extract_module
from ai_service import settings as ai_settings


def _font(size=34):
    for name in ("arial.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)


def _make_scanned_pdf(path):
    """A PDF that holds only a picture (no text layer): a heading, a sentence and a small table."""
    img = Image.new("RGB", (1500, 900), "white")
    d = ImageDraw.Draw(img)
    font = _font()
    d.text((80, 60), "Chapter 5. System Testing", fill="black", font=font)
    d.text((80, 140), "Unit tests were run before every release.", fill="black", font=font)
    rows = [("Test case", "Expected", "Result"), ("Login wrong password", "Error message", "Pass")]
    for r, row in enumerate(rows):
        y = 280 + r * 90
        for c, cell in enumerate(row):
            d.text((80 + c * 450, y), cell, fill="black", font=font)
    img.save(path, resolution=200)


def test_boxes_to_text_keeps_table_cells_apart():
    items = [
        (100, 100, 400, 140, "Heading line"),
        (100, 200, 250, 240, "Test case"),
        (600, 202, 760, 242, "Expected"),
        (100, 300, 380, 340, "Login failure"),
        (600, 301, 820, 341, "Error message"),
    ]
    text = extract_module._boxes_to_text(items)
    assert text.splitlines() == [
        "Heading line",
        "Test case | Expected",
        "Login failure | Error message",
    ]


def test_boxes_to_text_joins_close_words_with_a_space():
    items = [(100, 100, 200, 140, "Hello"), (215, 101, 330, 141, "world")]
    assert extract_module._boxes_to_text(items) == "Hello world"


def test_scanned_pdf_is_read_with_ocr(tmp_path):
    pytest.importorskip("rapidocr")
    path = tmp_path / "scan.pdf"
    _make_scanned_pdf(path)

    doc = extract_document(str(path))
    text = doc.pages[0].text

    assert "System Testing" in text
    assert any(
        "Login" in line and "Error message" in line and "Pass" in line and "|" in line
        for line in text.splitlines()
    )
    assert doc.pages[0].chapter and "Chapter 5" in doc.pages[0].chapter
    assert any("OCR" in w for w in doc.warnings)


def test_scanned_pdf_with_ocr_off_gives_the_scanned_warning(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    _make_scanned_pdf(path)
    monkeypatch.setattr(ai_settings, "OCR_ENABLED", False)

    doc = extract_document(str(path))

    assert not doc.pages[0].text
    assert any("scanned PDF" in w for w in doc.warnings)


def test_docx_nested_table_is_read(tmp_path):
    d = docx.Document()
    outer = d.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = "Outer cell"
    inner = outer.cell(0, 1).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "Inner left"
    inner.cell(0, 1).text = "Inner right"
    path = tmp_path / "nested.docx"
    d.save(path)

    text = "\n".join(p.text for p in extract_document(str(path)).pages)
    assert "Outer cell" in text
    assert "Inner left | Inner right" in text      

    # ---------- pictures inside DOCX files ----------

from docx.shared import Inches


def _make_picture(path, size=(1500, 600)):
       """A picture of text: a heading line and a small two-column table."""
       img = Image.new("RGB", size, "white")
       d = ImageDraw.Draw(img)
       font = _font()
       d.text((60, 40), "Figure 3. System Testing Summary", fill="black", font=font)
       rows = [("Test case", "Result"), ("Login wrong password", "Pass")]
       for r, row in enumerate(rows):
           for c, cell in enumerate(row):
               d.text((60 + c * 700, 200 + r * 100), cell, fill="black", font=font)
       img.save(path)


def test_docx_picture_text_is_read_with_ocr(tmp_path):
       pytest.importorskip("rapidocr")
       _make_picture(tmp_path / "pic.png")
       d = docx.Document()
       d.add_paragraph("Before the picture.")
       d.add_picture(str(tmp_path / "pic.png"), width=Inches(6))
       d.add_paragraph("After the picture.")
       path = tmp_path / "pic.docx"
       d.save(path)

       doc = extract_document(str(path))
       text = "\n".join(p.text for p in doc.pages)

       assert "System Testing Summary" in text
       assert any("Login" in line and "Pass" in line and "|" in line for line in text.splitlines())
       assert text.index("Before the picture.") < text.index("System Testing") < text.index("After the picture.")
       assert any("picture(s) in the DOCX were read with OCR" in w for w in doc.warnings)


def test_docx_picture_inside_a_table_cell_is_read(tmp_path):
       pytest.importorskip("rapidocr")
       _make_picture(tmp_path / "pic.png")
       d = docx.Document()
       table = d.add_table(rows=1, cols=2)
       table.cell(0, 0).text = "Screenshot"
       table.cell(0, 1).paragraphs[0].add_run().add_picture(str(tmp_path / "pic.png"), width=Inches(3))
       path = tmp_path / "cellpic.docx"
       d.save(path)

       text = "\n".join(p.text for p in extract_document(str(path)).pages)
       assert "Screenshot" in text
       assert "System Testing Summary" in text


def test_docx_tiny_picture_is_ignored(tmp_path):
       Image.new("RGB", (30, 30), "black").save(tmp_path / "icon.png")
       d = docx.Document()
       d.add_paragraph("Some text here.")
       d.add_picture(str(tmp_path / "icon.png"))
       path = tmp_path / "icon.docx"
       d.save(path)

       doc = extract_document(str(path))
       assert "Some text here." in doc.pages[0].text
       assert not any("OCR" in w for w in doc.warnings)


def test_docx_pictures_are_ignored_when_ocr_is_off(tmp_path, monkeypatch):
       _make_picture(tmp_path / "pic.png")
       d = docx.Document()
       d.add_paragraph("Only this text.")
       d.add_picture(str(tmp_path / "pic.png"), width=Inches(6))
       path = tmp_path / "off.docx"
       d.save(path)
       monkeypatch.setattr(ai_settings, "OCR_ENABLED", False)

       doc = extract_document(str(path))
       assert "System Testing" not in doc.pages[0].text
       assert not any("OCR" in w for w in doc.warnings)   