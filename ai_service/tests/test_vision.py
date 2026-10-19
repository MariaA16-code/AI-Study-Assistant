import docx
import pytest
from docx.shared import Inches
from PIL import Image, ImageDraw

from ai_service import extract as ex
from ai_service import settings, vision
from ai_service.extract import extract_document

FLOW = "Flowchart\nLogin -> Check password\nCheck password -> Dashboard"


@pytest.fixture
def vision_on(monkeypatch):
    monkeypatch.setattr(settings, "VISION_ENABLED", True)
    monkeypatch.setattr(settings, "GEMINI_API_KEY", "test-key")
    calls = []

    def fake(png, prompt):
        calls.append(prompt)
        return FLOW

    monkeypatch.setattr(vision, "_call_gemini", fake)
    return calls


def _picture(path):
    img = Image.new("RGB", (400, 300), "white")
    d = ImageDraw.Draw(img)
    d.rectangle((20, 20, 150, 80), outline="black")
    d.rectangle((220, 200, 380, 260), outline="black")
    d.line((150, 50, 300, 200), fill="black", width=3)
    img.save(path)


def _flowchart_pdf(path, boxes=5):
    canvas = pytest.importorskip("reportlab.pdfgen.canvas")
    c = canvas.Canvas(str(path))
    c.drawString(72, 800, "Chapter 4 Design")
    c.drawString(72, 780, "The login flow is shown below in a diagram.")
    for i in range(boxes):
        c.rect(72, 700 - i * 60, 150, 40)
    c.showPage()
    c.drawString(72, 780, "This page has only ordinary text and no figure at all.")
    c.save()


# ---- vision.py -----------------------------------------------------------------

def test_off_by_default_in_tests():
    assert vision.describe_image(Image.new("RGB", (10, 10))) is None


def test_none_reply_means_no_information(vision_on, monkeypatch):
    monkeypatch.setattr(vision, "_call_gemini", lambda png, prompt: "NONE")
    assert vision.describe_image(Image.new("RGB", (10, 10))) == ""


def test_failure_returns_none(vision_on, monkeypatch):
    def boom(png, prompt):
        raise RuntimeError("bad request")

    monkeypatch.setattr(vision, "_call_gemini", boom)
    assert vision.describe_image(Image.new("RGB", (10, 10))) is None


def test_long_description_is_cut(vision_on, monkeypatch):
    monkeypatch.setattr(vision, "_call_gemini", lambda png, prompt: "x" * 5000)
    assert len(vision.describe_image(Image.new("RGB", (10, 10)))) == vision.MAX_CHARS


def test_quota_error_stops_the_rest(vision_on, monkeypatch):
    monkeypatch.setattr(settings, "VISION_WORKERS", 1)
    calls = []

    def quota(png, prompt):
        calls.append(1)
        raise RuntimeError("429 RESOURCE_EXHAUSTED quota")

    monkeypatch.setattr(vision, "_call_gemini", quota)
    out = vision.describe_many([Image.new("RGB", (10, 10))] * 4)
    assert out == [None, None, None, None]
    assert len(calls) == 1

def test_fallback_model_is_tried_when_the_main_one_is_busy(vision_on, monkeypatch):
    monkeypatch.setattr(settings, "VISION_FALLBACK_MODEL", "backup-model")
    used = []

    def fake(png, prompt, model=None):
        used.append(model)
        if model is None:
            raise RuntimeError("503 UNAVAILABLE high demand")
        return "A -> B"

    monkeypatch.setattr(vision, "_call_gemini", fake)
    assert vision.describe_image(Image.new("RGB", (10, 10))) == "A -> B"
    assert used == [None, "backup-model"]


def test_busy_server_stops_the_rest(vision_on, monkeypatch):
    monkeypatch.setattr(settings, "VISION_WORKERS", 1)
    calls = []

    def busy(png, prompt):
        calls.append(1)
        raise RuntimeError("503 UNAVAILABLE high demand")

    monkeypatch.setattr(vision, "_call_gemini", busy)
    assert vision.describe_many([Image.new("RGB", (10, 10))] * 4) == [None] * 4
    assert len(calls) == 1  


def test_describe_many_keeps_order(vision_on, monkeypatch):
    seen = iter(["a", "b", "c"])
    monkeypatch.setattr(settings, "VISION_WORKERS", 1)
    monkeypatch.setattr(vision, "_call_gemini", lambda png, prompt: next(seen))
    assert vision.describe_many([Image.new("RGB", (10, 10))] * 3) == ["a", "b", "c"]


def test_figure_lines_all_carry_the_mark():
    assert vision.figure_lines("A -> B\n\n B -> C ") == ["[Figure] A -> B", "[Figure] B -> C"]


# ---- PDF -----------------------------------------------------------------------

def test_pdf_flowchart_page_is_described(tmp_path, vision_on):
    path = tmp_path / "flow.pdf"
    _flowchart_pdf(path)
    doc = extract_document(str(path))
    assert "[Figure] Login -> Check password" in doc.pages[0].text
    assert "ordinary text" in doc.pages[1].text and "[Figure]" not in doc.pages[1].text
    assert len(vision_on) == 1  # only the diagram page was sent
    assert vision_on[0] == vision.PAGE_PROMPT
    assert any("described with AI vision" in w for w in doc.warnings)
    assert doc.pages[0].chapter and "Chapter 4" in doc.pages[0].chapter


def test_pdf_text_only_makes_no_vision_call(tmp_path, vision_on):
    canvas = pytest.importorskip("reportlab.pdfgen.canvas")
    path = tmp_path / "text.pdf"
    c = canvas.Canvas(str(path))
    c.drawString(72, 780, "Plain text page about photosynthesis in green plants.")
    c.save()
    doc = extract_document(str(path))
    assert vision_on == []
    assert "[Figure]" not in doc.pages[0].text


def test_pdf_vision_off_leaves_text_unchanged(tmp_path):
    path = tmp_path / "flow.pdf"
    _flowchart_pdf(path)
    doc = extract_document(str(path))
    assert "[Figure]" not in doc.pages[0].text


def test_pdf_failure_gives_warning_and_keeps_text(tmp_path, vision_on, monkeypatch):
    def boom(png, prompt):
        raise RuntimeError("bad request")

    monkeypatch.setattr(vision, "_call_gemini", boom)
    path = tmp_path / "flow.pdf"
    _flowchart_pdf(path)
    doc = extract_document(str(path))
    assert "login flow" in doc.pages[0].text
    assert any("could not describe" in w for w in doc.warnings)


def test_pdf_page_cap(tmp_path, vision_on, monkeypatch):
    canvas = pytest.importorskip("reportlab.pdfgen.canvas")
    monkeypatch.setattr(settings, "VISION_MAX_PAGES", 1)
    path = tmp_path / "many.pdf"
    c = canvas.Canvas(str(path))
    for _ in range(3):
        c.drawString(72, 780, "A page with a diagram drawn below it.")
        for i in range(5):
            c.rect(72, 700 - i * 60, 150, 40)
        c.showPage()
    c.save()
    doc = extract_document(str(path))
    assert len(vision_on) == 1
    assert any("limited to 1 pages" in w and "2 page(s)" in w for w in doc.warnings)


def test_pdf_repeated_logo_is_not_a_figure(tmp_path, vision_on):
    canvas = pytest.importorskip("reportlab.pdfgen.canvas")
    logo = tmp_path / "logo.png"
    Image.new("RGB", (200, 200), "blue").save(logo)
    path = tmp_path / "logo.pdf"
    c = canvas.Canvas(str(path))
    for n in range(4):
        c.drawImage(str(logo), 72, 650, width=150, height=150)
        c.drawString(72, 600, f"Ordinary text on page {n + 1} of the document.")
        c.showPage()
    c.save()
    extract_document(str(path))
    assert vision_on == []


def test_pdf_embedded_picture_is_a_figure(tmp_path, vision_on):
    canvas = pytest.importorskip("reportlab.pdfgen.canvas")
    pic = tmp_path / "pic.png"
    _picture(pic)
    path = tmp_path / "pic.pdf"
    c = canvas.Canvas(str(path))
    c.drawImage(str(pic), 72, 500, width=300, height=225)
    c.drawString(72, 450, "Figure 1 shows the system overview in the picture above.")
    c.save()
    doc = extract_document(str(path))
    assert "[Figure] Login -> Check password" in doc.pages[0].text


# ---- DOCX ----------------------------------------------------------------------

def _docx_with_picture(tmp_path):
    _picture(tmp_path / "pic.png")
    d = docx.Document()
    d.add_paragraph("Before the picture.")
    d.add_picture(str(tmp_path / "pic.png"), width=Inches(5))
    d.add_paragraph("After the picture.")
    path = tmp_path / "pic.docx"
    d.save(path)
    return path


def test_docx_picture_is_described_in_order(tmp_path, vision_on):
    doc = extract_document(str(_docx_with_picture(tmp_path)))
    text = "\n".join(p.text for p in doc.pages)
    assert text.index("Before the picture.") < text.index("[Figure] Login -> Check password") < text.index("After the picture.")
    assert vision_on == [vision.PICTURE_PROMPT]
    assert any("described with AI vision" in w for w in doc.warnings)


def test_docx_same_picture_is_described_once(tmp_path, vision_on):
    _picture(tmp_path / "pic.png")
    d = docx.Document()
    d.add_picture(str(tmp_path / "pic.png"), width=Inches(4))
    d.add_picture(str(tmp_path / "pic.png"), width=Inches(4))
    path = tmp_path / "two.docx"
    d.save(path)
    doc = extract_document(str(path))
    assert len(vision_on) == 1
    assert "\n".join(p.text for p in doc.pages).count("[Figure] Login") == 2  # shown at both places


def test_docx_vision_failure_falls_back_to_ocr(tmp_path, vision_on, monkeypatch):
    def boom(png, prompt):
        raise RuntimeError("bad request")

    monkeypatch.setattr(vision, "_call_gemini", boom)
    monkeypatch.setattr(ex, "_ocr_image", lambda image: "Login\nDashboard")
    doc = extract_document(str(_docx_with_picture(tmp_path)))
    text = "\n".join(p.text for p in doc.pages)
    assert "Dashboard" in text and "[Figure]" not in text
    assert any("could not describe" in w for w in doc.warnings)


def test_docx_decorative_picture_adds_nothing(tmp_path, vision_on, monkeypatch):
    monkeypatch.setattr(vision, "_call_gemini", lambda png, prompt: "NONE")
    monkeypatch.setattr(ex, "_ocr_image", lambda image: pytest.fail("OCR must not run"))
    doc = extract_document(str(_docx_with_picture(tmp_path)))
    assert "[Figure]" not in "\n".join(p.text for p in doc.pages)


def test_docx_picture_cap_sends_the_rest_to_ocr(tmp_path, vision_on, monkeypatch):
    monkeypatch.setattr(settings, "VISION_MAX_IMAGES", 1)
    monkeypatch.setattr(ex, "_ocr_image", lambda image: "OCR words")
    d = docx.Document()
    for color in ("white", "yellow"):
        Image.new("RGB", (300, 200), color).save(tmp_path / f"{color}.png")
        d.add_picture(str(tmp_path / f"{color}.png"), width=Inches(3))
    path = tmp_path / "cap.docx"
    d.save(path)
    doc = extract_document(str(path))
    text = "\n".join(p.text for p in doc.pages)
    assert "[Figure] Login" in text and "OCR words" in text
    assert any("limited to 1 pictures" in w for w in doc.warnings)


def test_docx_vision_works_with_ocr_off(tmp_path, vision_on, monkeypatch):
    monkeypatch.setattr(settings, "OCR_ENABLED", False)
    doc = extract_document(str(_docx_with_picture(tmp_path)))
    assert "[Figure] Login -> Check password" in "\n".join(p.text for p in doc.pages)